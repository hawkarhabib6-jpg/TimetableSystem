# -*- coding: utf-8 -*-
"""Read the source .docx and emit a structured, classified JSON document.

Every paragraph is converted from the Ali-K legacy encoding to Unicode and
tagged with a semantic role, so the layout stage can style by meaning
instead of by the ad-hoc colours the original used.
"""
import json, re, sys, os
import docx
from docx.oxml.ns import qn
from alik import convert
from textfix import fill_marker_gaps
from edits import apply as apply_edits

SRC = os.path.join(os.path.dirname(__file__), '..', 'src', 'malzama.docx')
OUT = os.path.join(os.path.dirname(__file__), '..', 'build', 'doc.json')

KURDISH = re.compile(r'[؀-ۿ]')
LATIN = re.compile(r'[A-Za-z]')

UNIT_WORDS = ('one two three four five six seven eight nine ten eleven twelve '
              'thirteen fourteen fifteen sixteen').split()
UNIT_RE = re.compile(r'^\s*unit\s+(' + '|'.join(UNIT_WORDS) + r'|\d{1,2})\b', re.I)
TB_RE = re.compile(r'^\s*T\.?\s*B\b', re.I)
EG_RE = re.compile(r'^\s*e\.?\s*g\b', re.I)
MCQ_RE = re.compile(r'\(?\s*[aA][.)]\s.{0,90}?\b[bB][.)]\s', re.S)
OPT_RE = re.compile(r'^\s*\(?\s*([a-d])[.)]\s*(.+)$', re.S)

# Colours the author used as an informal semantic layer.
ROLE_BY_COLOUR = {
    '7030A0': 'kurdish', '7030a0': 'kurdish',
    'C00000': 'answer', 'FF0000': 'answer',
    '0070C0': 'note', '365F91': 'note', 'E36C0A': 'answer',
}


def script_of(t):
    k = len(KURDISH.findall(t))
    l = len(LATIN.findall(t))
    if k and k >= l:
        return 'ku'
    if l:
        return 'en'
    return 'neutral'


def para_runs(p):
    out = []
    for r in p.runs:
        if not r.text:
            continue
        col = None
        try:
            if r.font.color is not None and r.font.color.rgb is not None:
                col = str(r.font.color.rgb).upper()
        except Exception:
            pass
        hl = r._element.find('.//' + qn('w:highlight'))
        out.append({
            'text': convert(r.text),
            'b': bool(r.bold), 'i': bool(r.italic), 'u': bool(r.underline),
            'sz': r.font.size.pt if r.font.size else None,
            'color': col,
            'hl': hl.get(qn('w:val')) if hl is not None else None,
        })
    return out


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def textbox_text(p):
    """Text of any shape anchored in this paragraph.

    Word stores each shape twice — once as DrawingML, once as a VML fallback —
    so identical consecutive chunks are collapsed.
    """
    chunks = []
    for tx in p._element.iter(W + 'txbxContent'):
        parts = []
        for para in tx.iter(W + 'p'):
            t = ''.join(x.text or '' for x in para.iter(W + 't')).strip()
            if t:
                parts.append(t)
        t = convert(' '.join(parts)).strip()
        if t and (not chunks or chunks[-1] != t):
            chunks.append(t)
    return chunks


def images_in(p, rel_map):
    found = []
    for blip in p._element.findall('.//' + qn('a:blip')):
        rid = blip.get(qn('r:embed')) or blip.get(qn('r:link'))
        if rid in rel_map:
            found.append(rel_map[rid])
    VML = '{urn:schemas-microsoft-com:vml}imagedata'
    R_ID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    for im in p._element.findall('.//' + VML):
        rid = im.get(R_ID)
        if rid in rel_map:
            found.append(rel_map[rid])
    return found


def classify(text, runs, style, align):
    t = text.strip()
    if not t:
        return 'blank'
    low = t.lower()
    colours = {r['color'] for r in runs if r['color']}
    bold = any(r['b'] for r in runs)
    size = max([r['sz'] for r in runs if r['sz']] or [0])
    script = script_of(t)

    if UNIT_RE.match(t) and len(t) < 40:
        return 'unit'
    if TB_RE.match(t):
        return 'tb'
    if EG_RE.match(t):
        return 'example'
    if MCQ_RE.search(t):
        return 'mcq'
    if script == 'ku':
        return 'kurdish'
    for c in colours:
        if ROLE_BY_COLOUR.get(c) == 'answer':
            return 'answer'
    if style.startswith('Heading') or (bold and size >= 16 and len(t) < 70):
        return 'heading'
    if bold and len(t) < 80 and not t.endswith('.'):
        return 'subheading'
    if t.endswith('?'):
        return 'question'
    return 'text'


def main():
    d = docx.Document(SRC)
    rel_map = {}
    for rid, rel in d.part.rels.items():
        if 'image' in rel.reltype:
            rel_map[rid] = os.path.basename(rel.target_ref)

    blocks = []
    body = d.element.body
    pmap = {p._element: p for p in d.paragraphs}
    tmap = {t._element: t for t in d.tables}

    for el in body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == 'p':
            p = pmap.get(el)
            if p is None:
                continue
            runs = para_runs(p)
            text = convert(p.text).strip()
            imgs = images_in(p, rel_map)

            # Shapes carry the unit headings and most section titles.
            for box in textbox_text(p):
                kind = 'unit' if UNIT_RE.match(box) and len(box) < 40 else 'boxed'
                blocks.append({'kind': kind, 'text': box, 'runs': [],
                               'script': script_of(box), 'style': 'Shape',
                               'align': None, 'images': []})

            if not text and not imgs:
                continue
            blocks.append({
                'kind': classify(text, runs, p.style.name, p.alignment),
                'text': text,
                'runs': runs,
                'script': script_of(text),
                'style': p.style.name,
                'align': str(p.alignment).split(' ')[0] if p.alignment else None,
                'images': imgs,
            })
        elif tag == 'tbl':
            t = tmap.get(el)
            if t is None:
                continue
            rows = []
            for row in t.rows:
                rows.append([convert(c.text).strip() for c in row.cells])
            blocks.append({'kind': 'table', 'rows': rows,
                           'script': script_of(' '.join(' '.join(r) for r in rows))})

    gaps = fill_marker_gaps(blocks)
    edited = apply_edits(blocks)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(blocks, open(OUT, 'w', encoding='utf8'),
              ensure_ascii=False, indent=1)

    from collections import Counter
    c = Counter(b['kind'] for b in blocks)
    print('blocks:', len(blocks))
    for k, n in c.most_common():
        print(f'  {k:12s} {n}')
    print('images referenced:', sum(len(b.get("images", [])) for b in blocks))
    print('list markers the author left out:', gaps)
    print('editorial changes applied:', edited)


if __name__ == '__main__':
    main()
