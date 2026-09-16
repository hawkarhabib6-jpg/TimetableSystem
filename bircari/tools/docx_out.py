# -*- coding: utf-8 -*-
"""Write the booklet as an editable .docx.

Equations are real Word maths, built by latex2omml, so the reader can click
into one and change it. Kurdish runs right-to-left; the mathematics does not.
"""
import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

from latex2omml import omath

KU = 'Noto Naskh Arabic'
KUD = 'Noto Kufi Arabic'
EN = 'Cambria'
EN_SANS = 'Inter'

INK = RGBColor(0x17, 0x23, 0x3A)
KEY = RGBColor(0x1F, 0x5F, 0x8B)
LAW = RGBColor(0x8A, 0x5A, 0x00)
GOOD = RGBColor(0x17, 0x6B, 0x57)
WARN = RGBColor(0xA3, 0x3A, 0x2E)
SOFT = RGBColor(0x47, 0x56, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_KEY, HEX_KEY_WASH = '1F5F8B', 'EEF5FA'
HEX_LAW_WASH, HEX_LAW_EDGE = 'FFF8E8', 'EBD9AE'
HEX_GOOD_WASH, HEX_WARN_WASH = 'EDF6F3', 'FDF0EE'
HEX_RULE, HEX_ZEBRA = 'E3E6EC', 'FAFBFD'

PPR_SEQ = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
           'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd',
           'tabs', 'bidi', 'spacing', 'ind', 'contextualSpacing', 'jc',
           'textDirection', 'textAlignment', 'outlineLvl', 'rPr', 'sectPr']
TCPR_SEQ = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders',
            'shd', 'noWrap', 'tcMar', 'textDirection', 'vAlign', 'hideMark']


def _ins(parent, el):
    seq = TCPR_SEQ if parent.tag.endswith('tcPr') else PPR_SEQ
    name = el.tag.split('}')[1]
    if name not in seq:
        parent.append(el)
        return el
    rank = seq.index(name)
    for child in parent:
        cn = child.tag.split('}')[1]
        if cn in seq and seq.index(cn) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def _drop(parent, name):
    for c in parent.findall(qn('w:' + name)):
        parent.remove(c)


def shade(el, fill):
    _drop(el, 'shd')
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    _ins(el, s)


def borders(el, **sides):
    is_cell = el.tag.endswith('tcPr')
    name = 'tcBorders' if is_cell else 'pBdr'
    _drop(el, name)
    bd = OxmlElement('w:' + name)
    for side in ('top', 'left', 'bottom', 'right'):
        spec = sides.get(side)
        e = OxmlElement(f'w:{side}')
        if spec:
            sz, col = spec
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(sz))
            e.set(qn('w:color'), col)
        else:
            e.set(qn('w:val'), 'nil')
            e.set(qn('w:sz'), '0')
            e.set(qn('w:color'), 'auto')
        e.set(qn('w:space'), '4')
        bd.append(e)
    _ins(el, bd)


def rtl(par):
    pPr = par._p.get_or_add_pPr()
    _drop(pPr, 'bidi')
    _ins(pPr, OxmlElement('w:bidi'))
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in par.runs:
        rPr = r._element.get_or_add_rPr()
        if rPr.find(qn('w:rtl')) is None:
            rPr.append(OxmlElement('w:rtl'))


def spacing(par, before=0, after=4, line=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


# Every equation in the booklet is set in western digits. Prose that counts
# in Arabic-Indic ones beside them reads as a different number system on the
# same page, so the two are brought together here, in one place.
WESTERN = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')


def run(par, text, font=KU, size=11, bold=False, italic=False, color=None):
    r = par.add_run(text.translate(WESTERN))
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if color is not None:
        r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), font)
    return r


def math(par, tex):
    """Append one equation to a paragraph, as Word maths."""
    par._p.append(etree.fromstring(omath(tex)))


def cellpar(cell):
    ps = cell.paragraphs
    if len(ps) == 1 and not ps[0].runs and not ps[0].text:
        return ps[0]
    return cell.add_paragraph()


def onecell(doc, fill=None, right=None, edge=None, keep=False):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if keep:
        # A worked example is one row. Splitting it leaves its number on one
        # page and its working on the next.
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    c = t.cell(0, 0)
    c.paragraphs[0].text = ''
    if fill:
        shade(c._tc.get_or_add_tcPr(), fill)
    borders(c._tc.get_or_add_tcPr(), top=edge, bottom=edge, left=edge,
            right=right or edge)
    return t, c


# ------------------------------------------------------------- components
def heading_chapter(doc, no, title_ku, title_en):
    p = doc.add_paragraph()
    run(p, f'{no}  ', EN_SANS, 18, bold=True, color=KEY)
    run(p, title_ku, KUD, 15, bold=True, color=INK)
    rtl(p)
    borders(p._p.get_or_add_pPr(), bottom=(14, HEX_KEY))
    spacing(p, 14, 2)
    q = doc.add_paragraph()
    run(q, title_en, EN_SANS, 9.5, color=SOFT)
    spacing(q, 0, 8)
    _keep(p)


def heading_section(doc, title):
    p = doc.add_paragraph()
    run(p, title, KUD, 12.5, bold=True, color=KEY)
    rtl(p)
    borders(p._p.get_or_add_pPr(), bottom=(6, HEX_RULE))
    spacing(p, 10, 3)
    _keep(p)
    return p


def _keep(par):
    pPr = par._p.get_or_add_pPr()
    _drop(pPr, 'keepNext')
    _ins(pPr, OxmlElement('w:keepNext'))


def para(doc, text, size=11, color=None):
    p = doc.add_paragraph()
    run(p, text, KU, size, color=color)
    rtl(p)
    spacing(p, 0, 4, 1.4)
    return p


def concept(doc, text):
    _, c = onecell(doc, HEX_KEY_WASH, right=(18, HEX_KEY))
    p = cellpar(c)
    run(p, text, KU, 11, color=INK)
    rtl(p)
    spacing(p, 2, 2, 1.4)
    doc.add_paragraph()


def bullets(doc, items):
    for t in items:
        p = doc.add_paragraph()
        run(p, '•  ', EN_SANS, 10, color=KEY)
        if isinstance(t, tuple):          # (plain, tex, plain)
            for piece in t:
                if piece.startswith('$') and piece.endswith('$'):
                    math(p, piece[1:-1])
                else:
                    run(p, piece, KU, 11)
        else:
            run(p, t, KU, 11)
        rtl(p)
        p.paragraph_format.right_indent = Mm(4)
        spacing(p, 0, 2, 1.35)


def law(doc, title, lines, note=None):
    _, c = onecell(doc, HEX_LAW_WASH, edge=(8, HEX_LAW_EDGE))
    p = cellpar(c)
    run(p, title, KUD, 10.5, bold=True, color=LAW)
    rtl(p)
    spacing(p, 1, 3)
    for tex in lines:
        q = c.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math(q, tex)
        spacing(q, 2, 2)
    if note:
        n = c.add_paragraph()
        run(n, note, KU, 9.8, color=SOFT)
        rtl(n)
        borders(n._p.get_or_add_pPr(), top=(4, HEX_LAW_EDGE))
        spacing(n, 3, 1, 1.35)
    doc.add_paragraph()


def example(doc, no, prompt_parts, figure=None, steps=(), answer=None):
    _, c = onecell(doc, 'FFFFFF', edge=(6, HEX_RULE), keep=True)

    head = cellpar(c)
    run(head, f'نموونە {no}', KUD, 9.5, bold=True, color=WHITE)
    shade(head._p.get_or_add_pPr(), HEX_KEY)
    rtl(head)
    spacing(head, 1, 2)

    if prompt_parts:
        p = c.add_paragraph()
        # A question that is only an equation sits centred, like the working
        # below it; one with words in it reads from the right.
        bare = all(x.startswith('$') and x.endswith('$') for x in prompt_parts)
        for piece in prompt_parts:
            if piece.startswith('$') and piece.endswith('$'):
                math(p, piece[1:-1])
            else:
                run(p, piece, KU, 11)
        if bare:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            rtl(p)
        spacing(p, 1, 3, 1.4)

    if figure:
        f = c.add_paragraph()
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            f.add_run().add_picture(figure, width=Mm(62))
        except Exception:
            pass
        spacing(f, 2, 3)

    for i, (why, tex) in enumerate(steps, 1):
        # A step may be equation-only, continuing the one above it.
        if why is not None:
            w = c.add_paragraph()
            run(w, f'{i}. ', EN_SANS, 9.5, bold=True, color=KEY)
            for piece in (why if isinstance(why, tuple) else (why,)):
                if piece.startswith('$') and piece.endswith('$'):
                    math(w, piece[1:-1])
                else:
                    run(w, piece, KU, 10.6, color=SOFT)
            rtl(w)
            spacing(w, 2, 1, 1.35)
        if tex:
            e = c.add_paragraph()
            e.alignment = WD_ALIGN_PARAGRAPH.CENTER
            math(e, tex)
            spacing(e, 1, 2)

    if answer:
        a = c.add_paragraph()
        run(a, 'وەڵام:  ', KUD, 10.5, bold=True, color=GOOD)
        if answer.startswith('$'):
            math(a, answer[1:-1])
        else:
            run(a, answer, KU, 11, bold=True, color=GOOD)
        shade(a._p.get_or_add_pPr(), HEX_GOOD_WASH)
        rtl(a)
        spacing(a, 3, 2)
    doc.add_paragraph()


def callout(doc, title, items, fill, colour, edge):
    _, c = onecell(doc, fill, right=(18, edge))
    p = cellpar(c)
    run(p, title, KUD, 10.4, bold=True, color=colour)
    rtl(p)
    spacing(p, 1, 2)
    for t in items:
        q = c.add_paragraph()
        run(q, '– ', EN_SANS, 10, color=colour)
        for piece in (t if isinstance(t, tuple) else (t,)):
            if piece.startswith('$') and piece.endswith('$'):
                math(q, piece[1:-1])
            else:
                run(q, piece, KU, 10.5)
        rtl(q)
        spacing(q, 0, 2, 1.35)
    doc.add_paragraph()


def mistakes(doc, items):
    callout(doc, 'هەڵە باوەکان', items, HEX_WARN_WASH, WARN, 'A33A2E')


def tips(doc, items):
    callout(doc, 'ڕێگا کورتەکان', items, HEX_GOOD_WASH, GOOD, '176B57')


def table(doc, headers, rows):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Without this the first column lands on the left, so a right-to-left
    # table reads back to front.
    tblPr = t._tbl.tblPr
    vis = OxmlElement('w:bidiVisual')
    tblPr.append(vis)
    for ci, h in enumerate(headers):
        cell = t.cell(0, ci)
        p = cell.paragraphs[0]
        run(p, h, KUD, 9.6, bold=True, color=WHITE)
        rtl(p)
        spacing(p, 1, 1)
        shade(cell._tc.get_or_add_tcPr(), HEX_KEY)
    for ri, r in enumerate(rows, 1):
        for ci, val in enumerate(r):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            for piece in (val if isinstance(val, tuple) else (val,)):
                if piece.startswith('$') and piece.endswith('$'):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    math(p, piece[1:-1])
                else:
                    run(p, piece, KU, 10.4)
                    rtl(p)
            spacing(p, 1, 1)
            if ri % 2 == 0:
                shade(cell._tc.get_or_add_tcPr(), HEX_ZEBRA)
    doc.add_paragraph()
