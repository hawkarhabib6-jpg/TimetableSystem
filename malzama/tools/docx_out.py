# -*- coding: utf-8 -*-
"""Emit the booklet as an editable .docx carrying the same design system.

Full-page graphics (cover, part dividers, unit openers) are placed as images
rendered from the same HTML the PDF uses, so both outputs stay identical.
Everything else is real Word content: styled paragraphs, tables and pictures
the user can edit.
"""
import json, os, re, sys
import docx
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import book
from mixed import segments
from textfix import clean, split_marker
from render import WORD2NUM, UNIT_TITLES, split_options, FORMULA_HINT

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
IMG = os.path.join(ROOT, 'assets', 'img')
GFX = os.path.join(ROOT, 'build', 'gfx')

EN = 'Source Serif 4'
EN_SANS = 'Inter'
KU = 'Noto Naskh Arabic'
KU_DISP = 'Noto Kufi Arabic'

INK = RGBColor(0x12, 0x20, 0x3A)
KU_C = RGBColor(0x5F, 0x3F, 0xA8)
ANS = RGBColor(0x0E, 0x7C, 0x66)
TBC = RGBColor(0x4A, 0x3A, 0x1C)
SUN = RGBColor(0xD4, 0x44, 0x2A)


# ---------------------------------------------------------------- low level
# OOXML fixes the order of children inside pPr and tcPr; appending out of
# sequence produces a file Word and LibreOffice both refuse to open.
PPR_SEQ = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
           'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd',
           'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
           'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN',
           'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
           'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
           'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl',
           'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
TCPR_SEQ = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders',
            'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign',
            'hideMark', 'headers', 'cellIns', 'cellDel', 'cellMerge',
            'tcPrChange']


def insert_ordered(parent, el):
    seq = TCPR_SEQ if parent.tag.endswith('tcPr') else PPR_SEQ
    name = el.tag.split('}')[1]
    if name not in seq:
        parent.append(el)
        return el
    rank = seq.index(name)
    for child in parent:
        cname = child.tag.split('}')[1]
        if cname in seq and seq.index(cname) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def drop(parent, name):
    for child in parent.findall(qn('w:' + name)):
        parent.remove(child)


def shade(el, fill):
    drop(el, 'shd')
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    insert_ordered(el, sh)


def cell_shade(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def borders(el, **sides):
    """sides: top/bottom/left/right = (size_eighths, colour) or None."""
    is_cell = el.tag.endswith('tcPr')
    name = 'tcBorders' if is_cell else 'pBdr'
    drop(el, name)
    bd = OxmlElement('w:' + name)
    for side in ('top', 'left', 'bottom', 'right'):
        spec = sides.get(side)
        e = OxmlElement(f'w:{side}')
        if spec:
            sz, col = spec
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(sz))
            e.set(qn('w:color'), col)
        else:
            e.set(qn('w:val'), 'nil')
            e.set(qn('w:sz'), '0')
            e.set(qn('w:color'), 'auto')
        e.set(qn('w:space'), '4')
        bd.append(e)
    insert_ordered(el, bd)


def rtl(par):
    pPr = par._p.get_or_add_pPr()
    drop(pPr, 'bidi')
    insert_ordered(pPr, OxmlElement('w:bidi'))
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in par.runs:
        rPr = r._element.get_or_add_rPr()
        if rPr.find(qn('w:rtl')) is None:
            rPr.append(OxmlElement('w:rtl'))


def spacing(par, before=0, after=3, line=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def run(par, text, font=EN, size=10, bold=False, italic=False, color=None):
    r = par.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), font)
    return r


def rich(par, text, host, size=10, bold=False, italic=False, color=None,
         ku_font=None, en_font=None):
    """Write text as one run per script, each in its own font.

    A bracket belongs to whatever it encloses, so the runs come from
    mixed.segments rather than from a naive split on character ranges.
    """
    ku_font = ku_font or KU
    en_font = en_font or EN
    made = []
    for sc, chunk in segments(clean(text)):
        font = ku_font if sc == 'ku' else en_font
        sz = size * (1.06 if sc == 'ku' and host == 'en' else 1)
        r = run(par, chunk, font, sz, bold=bold, italic=italic, color=color)
        if sc == 'ku':
            rPr = r._element.get_or_add_rPr()
            rPr.append(OxmlElement('w:rtl'))
        made.append(r)
    return made


def keep_next(par):
    """Word's own 'keep with next', so a heading is never left page-bottom."""
    pPr = par._p.get_or_add_pPr()
    drop(pPr, 'keepNext')
    insert_ordered(pPr, OxmlElement('w:keepNext'))


def onecell(doc, fill=None, left=None, edge=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    c.paragraphs[0].text = ''
    if fill:
        cell_shade(c, fill)
    borders(c._tc.get_or_add_tcPr(),
            top=edge, bottom=edge, right=edge,
            left=left or edge)
    return t, c


def cellpar(cell, first=[True]):
    """First call reuses the cell's empty paragraph; later ones append."""
    ps = cell.paragraphs
    if len(ps) == 1 and not ps[0].runs and not ps[0].text:
        return ps[0]
    return cell.add_paragraph()


# ---------------------------------------------------------------- sections
def set_margins(sec, top, bottom, left, right):
    sec.top_margin, sec.bottom_margin = Mm(top), Mm(bottom)
    sec.left_margin, sec.right_margin = Mm(left), Mm(right)


HEAD_TITLE = 'Sunrise 12'
HEAD_AUTHOR = 'Falah H. Younis'

FOOT_LEFT = 'Ibrahim Ahmad preparatory school'
FOOT_CENTRE = 'پەیمانگای ژیر'
FOOT_RIGHT = 'Shahid Aram preparatory school'


def clear_style_tabs(par):
    """Drop the tab stops the Header and Footer styles bring with them.

    Both built-in styles define their own centre and right tabs; left in
    place, a tab lands on whichever comes first and the text never reaches
    the margin.
    """
    pPr = par._p.get_or_add_pPr()
    drop(pPr, 'pStyle')
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '0')
    ind.set(qn('w:right'), '0')
    ind.set(qn('w:firstLine'), '0')
    drop(pPr, 'ind')
    insert_ordered(pPr, ind)


def tab_stops(par, positions):
    pPr = par._p.get_or_add_pPr()
    drop(pPr, 'tabs')
    tabs = OxmlElement('w:tabs')
    for pos, align in positions:
        t = OxmlElement('w:tab')
        t.set(qn('w:val'), align)
        t.set(qn('w:pos'), str(int(pos)))   # twentieths of a point
        tabs.append(t)
    insert_ordered(pPr, tabs)


def set_footer(section, on):
    """School line across the foot of every text page."""
    section.footer.is_linked_to_previous = False
    f = section.footer
    for extra in f.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p = f.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.text = ''
    if not on:
        return
    width = Mm(182).twips
    clear_style_tabs(p)
    tab_stops(p, [(width / 2, 'center'), (width, 'right')])
    grey = RGBColor(0x7A, 0x86, 0x9F)
    run(p, FOOT_LEFT, EN_SANS, 6.6, color=grey)
    run(p, '\t', EN_SANS, 6.6)
    r = run(p, FOOT_CENTRE, KU_DISP, 7.4, color=KU_C)
    r._element.get_or_add_rPr().append(OxmlElement('w:rtl'))
    run(p, '\t', EN_SANS, 6.6)
    run(p, FOOT_RIGHT, EN_SANS, 6.6, color=grey)
    borders(p._p.get_or_add_pPr(), top=(4, 'D8D2C8'))
    spacing(p, 2, 0)


def drop_trailing_empties(doc):
    """Remove spacer paragraphs left at the end of the body.

    Components leave an empty paragraph behind for spacing; one sitting just
    before a section break becomes a blank page of its own.
    """
    body = doc.element.body
    for el in reversed(list(body.iterchildren())):
        if not el.tag.endswith('}p'):
            break
        if ''.join(el.itertext()).strip() or el.findall('.//' + qn('w:drawing')):
            break
        body.remove(el)


def paragraphs_of(body):
    return [el for el in body.iterchildren() if el.tag.endswith('}p')]


def new_section(doc):
    """Start a section without leaving an empty paragraph behind.

    python-docx parks the outgoing section's properties in a paragraph of
    their own, which prints as a blank page. Word keeps them on the last
    paragraph of the section instead, so they are moved there.
    """
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    body = doc.element.body
    ps = paragraphs_of(body)
    if len(ps) < 2:
        return sec
    holder = ps[-1]
    pPr = holder.find(qn('w:pPr'))
    sectPr = pPr.find(qn('w:sectPr')) if pPr is not None else None
    prev = holder.getprevious()
    if sectPr is None or prev is None or not prev.tag.endswith('}p'):
        return sec
    prevPr = prev.find(qn('w:pPr'))
    if prevPr is None:
        prevPr = OxmlElement('w:pPr')
        prev.insert(0, prevPr)
    prevPr.append(sectPr)
    body.remove(holder)
    return sec


BLEED = (0, 0, 0, 0)
BODY = (18, 15, 14, 14)


def ensure_mode(doc, mode):
    """Switch page geometry only when it actually changes.

    Artwork pages run to the paper edge, text pages carry margins. Breaking
    the section between two consecutive artwork pages would leave a stray
    page between them, so the break happens only at a real transition.
    """
    if getattr(ensure_mode, 'cur', None) == mode:
        return
    drop_trailing_empties(doc)
    sec = (doc.sections[0] if not paragraphs_of(doc.element.body)
           else new_section(doc))
    set_margins(sec, *(BLEED if mode == 'bleed' else BODY))
    ensure_mode.cur = mode


def set_header(section, on):
    """Book title and author across the head of every text page."""
    section.header.is_linked_to_previous = False
    h = section.header
    for extra in h.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p = h.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.text = ''
    if not on:
        return
    clear_style_tabs(p)
    tab_stops(p, [(Mm(182).twips, 'right')])
    r = run(p, HEAD_TITLE.upper(), EN_SANS, 7.4, bold=True, color=INK)
    r.font.all_caps = True
    run(p, '\t', EN_SANS, 7)
    run(p, HEAD_AUTHOR, EN_SANS, 7, color=RGBColor(0x7A, 0x86, 0x9F))
    borders(p._p.get_or_add_pPr(), bottom=(4, 'D8D2C8'))
    spacing(p, 0, 3)


def fullpage(doc, path):
    """Place one full-bleed page image."""
    ensure_mode(doc, 'bleed')
    p = doc.add_paragraph()
    spacing(p, 0, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Mm(210), height=Mm(297))


# ---------------------------------------------------------------- blocks
def add_ku(doc_or_cell, text, size=10.5, color=KU_C, indent=True):
    p = (cellpar(doc_or_cell) if hasattr(doc_or_cell, '_tc')
         else doc_or_cell.add_paragraph())
    marker, text = split_marker(text)
    if marker:
        # Written first and left-to-right, so the number opens the line
        # rather than landing wherever bidi resolves its digits.
        run(p, f'{marker}-  ', EN_SANS, size * .92, bold=True, color=KU_C)
    rich(p, text, 'ku', size, color=color)
    rtl(p)
    spacing(p, 0, 3, 1.45)
    if indent and not hasattr(doc_or_cell, '_tc'):
        p.paragraph_format.right_indent = Mm(2)
        borders(p._p.get_or_add_pPr(), right=(18, '8A6FC7'))
        shade(p._p.get_or_add_pPr(), 'F7F4FD')
    return p


def add_en(doc, text, size=10, bold=False, color=None, font=EN):
    p = doc.add_paragraph()
    rich(p, text, 'en', size, bold=bold, color=color, en_font=font)
    spacing(p, 0, 3)
    return p


def add_mcq(cell_or_doc, opts):
    host = cell_or_doc
    wide = max(len(o[1] if isinstance(o, tuple) else o) for o in opts) > 46
    cols = 1 if wide else 2
    rows = (len(opts) + cols - 1) // cols
    t = host.add_table(rows=rows, cols=cols) if hasattr(host, 'add_table') \
        else host.add_table(rows=rows, cols=cols)
    if hasattr(t, 'style'):
        t.style = 'Table Grid'
    for i, o in enumerate(opts):
        letter, text = o if isinstance(o, tuple) else ('abcd'[i], o)
        c = t.cell(i // cols, i % cols)
        borders(c._tc.get_or_add_tcPr(), top=None, bottom=None,
                left=None, right=None)
        p = c.paragraphs[0]
        run(p, f'({letter})  ', EN_SANS, 8, bold=True, color=INK)
        rich(p, text, 'en', 9.5)
        spacing(p, 1, 1)
    return t


def build(blocks, units, doc):
    gfx = sorted(os.listdir(GFX)) if os.path.isdir(GFX) else []
    gi = 0

    def next_gfx():
        nonlocal gi
        if gi < len(gfx):
            path = os.path.join(GFX, gfx[gi])
            gi += 1
            fullpage(doc, path)

    drop_trailing_empties(doc)   # the body starts with one empty paragraph
    next_gfx()   # cover
    next_gfx()   # contents

    last_part, qn_ = None, 0
    for pi, n, s, e in units:
        if pi != last_part:
            next_gfx()          # part divider
            last_part = pi
        next_gfx()              # unit opener
        qn_ = 0

        ensure_mode(doc, 'body')
        grouped = book.render.group(blocks, s, e)

        def translation(at):
            """Every Kurdish block following `at`, and no more.

            A paragraph's translation is exactly the Kurdish that follows it:
            stopping at the first block would drop the rest of the same
            translation, and running on would absorb the next one."""
            got, j = [], at + 1
            while j < len(grouped) and grouped[j]['kind'] == 'kurdish' \
                    and not grouped[j].get('images'):
                got.append(grouped[j]['text'])
                j += 1
            return got, j - at - 1

        i = 0
        while i < len(grouped):
            b = grouped[i]
            k = b['kind']
            t = b.get('text', '')
            nxt = grouped[i + 1] if i + 1 < len(grouped) else None
            ku_all, ku_n = translation(i)
            ku = ku_all[0] if ku_all else None

            if k == 'unit' or (k == 'blank' and not b.get('images')):
                i += 1
                continue

            if b.get('images'):
                for name in b['images']:
                    if name.lower().endswith('.wmf'):
                        continue
                    path = os.path.join(IMG, name)
                    if not os.path.exists(path):
                        continue
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    try:
                        p.add_run().add_picture(path, width=Mm(120))
                        build.pics = getattr(build, 'pics', 0) + 1
                    except Exception:
                        build.skipped = getattr(build, 'skipped', 0) + 1
                    spacing(p, 4, 4)
                if t and len(t) > 40:
                    add_en(doc, t)
                i += 1
                continue

            if k == 'table':
                rows = b['rows']
                if not rows:
                    i += 1
                    continue
                tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
                tb.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    for ci, val in enumerate(r):
                        if ci >= len(rows[0]):
                            continue
                        c = tb.cell(ri, ci)
                        p = c.paragraphs[0]
                        is_ku = bool(re.search(r'[؀-ۿ]', val))
                        rich(p, val, 'ku' if is_ku else 'en', 9,
                             bold=(ri == 0),
                             color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0
                             else (KU_C if is_ku else None))
                        if is_ku:
                            rtl(p)
                        spacing(p, 1, 1)
                        cell_shade(c, '12203A' if ri == 0
                                   else ('FBF9F6' if ri % 2 == 0 else 'FFFFFF'))
                doc.add_paragraph()
                i += 1
                continue

            if k == 'formula':
                segs = [x.strip() for x in re.split(r'\s*\+\s*', t) if x.strip()]
                alts = b.get('alts', [])
                cells = []
                for idx, seg in enumerate(segs):
                    if idx == 0 and alts:
                        w = seg.split()
                        cells.append(w[0])
                        cells.append('\n'.join(dict.fromkeys(
                            [' '.join(w[1:])] + alts)))
                    else:
                        cells.append(seg)
                tb = doc.add_table(rows=1, cols=len(cells) * 2 - 1)
                tb.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ci in range(len(cells) * 2 - 1):
                    c = tb.cell(0, ci)
                    p = c.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if ci % 2:
                        run(p, '+', EN_SANS, 12, bold=True, color=INK)
                        borders(c._tc.get_or_add_tcPr())
                    else:
                        val = cells[ci // 2]
                        first = ci == 0
                        for li, line in enumerate(val.split('\n')):
                            pp = p if li == 0 else c.add_paragraph()
                            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            run(pp, line, EN_SANS, 9.5, bold=True,
                                color=RGBColor(0xFF, 0xFF, 0xFF) if first
                                else SUN)
                            spacing(pp, 1, 1)
                        cell_shade(c, '12203A' if first else 'FFFFFF')
                        borders(c._tc.get_or_add_tcPr(),
                                top=(6, 'E7E2DA'), bottom=(6, 'E7E2DA'),
                                left=(6, 'E7E2DA'), right=(6, 'E7E2DA'))
                doc.add_paragraph()
                i += 1
                continue

            if k == 'tb':
                t2 = re.sub(r'^\s*T\.?\s*B\s*\d*\s*[:/]?\s*', '', t, flags=re.I)
                _, c = onecell(doc, 'FFF7EA', edge=(6, 'F0DFC0'))
                p = cellpar(c)
                run(p, 'T.B  ', EN_SANS, 7.5, bold=True, color=RGBColor(0xA4, 0x64, 0x0B))
                is_ku = b.get('script') == 'ku'
                rich(p, t2, 'ku' if is_ku else 'en',
                     10 if is_ku else 9.5, color=TBC)
                if is_ku:
                    rtl(p)
                spacing(p, 1, 1)
                for x in ku_all:
                    add_ku(c, x, 10, TBC)
                i += ku_n
                doc.add_paragraph()
                i += 1
                continue

            if k == 'example':
                body = re.sub(r'^\s*e\.?\s*g\s*/?\s*', '', t, flags=re.I)
                _, c = onecell(doc, 'FFF4E6', left=(18, 'FFB13D'))
                p = cellpar(c)
                run(p, 'EXAMPLE', EN_SANS, 6.5, bold=True, color=SUN)
                spacing(p, 1, 1)
                p2 = c.add_paragraph()
                rich(p2, body, 'en', 9.5, italic=True,
                     color=RGBColor(0x5A, 0x32, 0x12))
                spacing(p2, 0, 1)
                doc.add_paragraph()
                i += 1
                continue

            if re.match(r'^\s*question\s*bank\b', t, re.I) and len(t) < 30:
                _, c = onecell(doc, '12203A')
                p = cellpar(c)
                run(p, 'Question Bank', EN_SANS, 11.5, bold=True,
                    color=RGBColor(0xFF, 0xFF, 0xFF))
                run(p, '                                        ', EN, 11)
                run(p, 'بانکی پرسیار', KU_DISP, 10.5, bold=True,
                    color=RGBColor(0xFF, 0xD9, 0xA8))
                spacing(p, 2, 2)
                keep_next(p)
                doc.add_paragraph()
                i += 1
                continue

            if k in ('heading', 'boxed'):
                p = add_en(doc, t, 14, bold=True, color=INK, font=EN_SANS)
                spacing(p, 10, 3)
                keep_next(p)
                i += 1
                continue

            if k == 'subheading':
                is_ku = b.get('script') == 'ku'
                p = doc.add_paragraph()
                rich(p, t, 'ku' if is_ku else 'en', 11, bold=True, color=INK,
                     ku_font=KU_DISP, en_font=EN_SANS)
                if is_ku:
                    rtl(p)
                borders(p._p.get_or_add_pPr(), bottom=(4, 'E7E2DA'))
                spacing(p, 6, 2)
                keep_next(p)
                i += 1
                continue

            if k in ('mcq', 'looseq'):
                if k == 'mcq':
                    parsed = split_options(t)
                    if not parsed:
                        add_en(doc, t)
                        i += 1
                        continue
                    lead, opts = parsed
                else:
                    lead, opts = t, b['opts']
                if not lead:
                    prev = grouped[i - 1] if i else None
                    if prev and prev['kind'] == 'text' and not prev.get('images') \
                            and prev.get('script') == 'en' \
                            and prev.get('text') and len(prev['text']) < 120 \
                            and not getattr(build, 'used_stem', None) == id(prev):
                        lead = prev['text']
                        build.used_stem = id(prev)
                        if doc.paragraphs and doc.paragraphs[-1].text == prev['text']:
                            pp = doc.paragraphs[-1]._p
                            pp.getparent().remove(pp)
                qn_ += 1
                _, c = onecell(doc, 'FFFFFF', left=(20, '1B5FA8'),
                               edge=(4, 'E7E2DA'))
                p = cellpar(c)
                run(p, f'{qn_}.  ', EN_SANS, 9.5, bold=True,
                    color=RGBColor(0x1B, 0x5F, 0xA8))
                rich(p, lead or 'Choose the correct answer', 'en', 10,
                     bold=True, color=INK, en_font=EN_SANS)
                spacing(p, 1, 2)
                add_mcq(c, opts)
                for x in ku_all:
                    add_ku(c, x, 10)
                i += ku_n
                doc.add_paragraph()
                i += 1
                continue

            if k == 'question':
                qn_ += 1
                _, c = onecell(doc, 'FFFFFF', left=(20, '1B5FA8'),
                               edge=(4, 'E7E2DA'))
                p = cellpar(c)
                run(p, f'{qn_}.  ', EN_SANS, 9.5, bold=True,
                    color=RGBColor(0x1B, 0x5F, 0xA8))
                rich(p, t, 'en', 10, bold=True, color=INK, en_font=EN_SANS)
                spacing(p, 1, 2)
                if nxt and nxt['kind'] == 'answer':
                    pa = c.add_paragraph()
                    run(pa, '✓  ', EN_SANS, 9, bold=True, color=ANS)
                    rich(pa, nxt['text'], 'en', 9.5, bold=True, color=ANS)
                    shade(pa._p.get_or_add_pPr(), 'EDF7F4')
                    spacing(pa, 1, 1)
                    i += 1
                    ku_all, ku_n = translation(i)
                for x in ku_all:
                    add_ku(c, x, 10)
                i += ku_n
                doc.add_paragraph()
                i += 1
                continue

            if k == 'kurdish':
                add_ku(doc, t, 10.5, indent=False)
                i += 1
                continue

            # text / answer
            red = any((r.get('color') or '') in ('C00000', 'FF0000', 'E36C0A')
                      for r in b.get('runs', []))
            is_prompt = bool(re.search(r'(:-|:|-|…|\.{3,}|_{3,})\s*$', t)) or \
                re.match(r'^\s*(find|choose|complete|correct|rewrite|put|match|'
                         r'answer|fill|write|underline|question)\b', t, re.I)
            if k == 'answer' and is_prompt and len(t) < 70:
                p = doc.add_paragraph()
                rich(p, t.rstrip(' -:'), 'en', 9.5, bold=True, color=SUN,
                     en_font=EN_SANS)
                borders(p._p.get_or_add_pPr(), bottom=(4, 'E7E2DA'))
                spacing(p, 5, 2)
                keep_next(p)
            elif k == 'answer' and red and len(t) > 25:
                p = doc.add_paragraph()
                run(p, '✓  ', EN_SANS, 9, bold=True, color=ANS)
                rich(p, t, 'en', 9.5, bold=True, color=ANS)
                shade(p._p.get_or_add_pPr(), 'EDF7F4')
                spacing(p, 2, 2)
            elif ku_all:
                add_en(doc, t)
                for x in ku_all:
                    add_ku(doc, x)
                i += ku_n
            else:
                add_en(doc, t)
            i += 1


def main():
    blocks = json.load(open(os.path.join(ROOT, 'build', 'doc.json'),
                            encoding='utf8'))
    units = book.segment(blocks)

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = EN
    st.font.size = Pt(10)
    set_margins(doc.sections[0], *BODY)
    doc.sections[0].page_width = Mm(210)
    doc.sections[0].page_height = Mm(297)

    build(blocks, units, doc)

    # Full-bleed sections run to the paper edge and carry no footer.
    for sec in doc.sections:
        text_page = bool(sec.left_margin and sec.left_margin > 0)
        set_header(sec, text_page)
        set_footer(sec, text_page)

    print('pictures placed:', getattr(build, 'pics', 0),
          '| unreadable:', getattr(build, 'skipped', 0))
    out = os.path.join(ROOT, 'build', 'Sunrise12-Companion.docx')
    doc.save(out)
    print('wrote', out, round(os.path.getsize(out) / 1e6, 1), 'MB')


if __name__ == '__main__':
    main()
